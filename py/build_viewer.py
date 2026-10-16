"""Turn the HRB 2026 producers workbook into a single self-contained HTML viewer.

The page is built to work in the iOS Files preview (Quick Look), where
JavaScript may not run: every view switch and filter is plain CSS driven by
radio buttons. When the page does get JavaScript (Safari, a desktop browser),
it adds search, "happening now" markers, remembered selections and a
staleness warning on top.

    python build_viewer.py workbook.xlsx out.html      # one-off build
"""

from __future__ import annotations

import colorsys
import datetime as dt
import html
import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

BUILD_VERSION = "1.0.0"

WEEKDAYS = ("monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday")
ROOM_COLORS = ("#7c3aed", "#ea580c", "#0d9488", "#2563eb", "#db2777", "#65a30d", "#ca8a04", "#0891b2")


# --------------------------------------------------------------------------- text + time helpers

def clean(value) -> str:
    if value is None:
        return ""
    s = str(value).replace("​", "").replace("﻿", "").replace("\xa0", " ").replace("\r", "")
    s = "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in s.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", s).strip()


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-") or "x"


def parse_time(value) -> dt.time | None:
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, dt.time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, (int, float)) and 0 <= value < 1:
        minutes = round(value * 1440)
        return dt.time(minutes // 60 % 24, minutes % 60)
    m = re.match(r"^\s*(\d{1,2})(?::(\d{2}))?(?::\d{2})?\s*([ap])?\.?\s*m?\.?\s*$", str(value), re.I)
    if not m:
        return None
    h, mi, ap = int(m.group(1)), int(m.group(2) or 0), (m.group(3) or "").lower()
    if ap == "p" and h < 12:
        h += 12
    if ap == "a" and h == 12:
        h = 0
    if h > 23 or mi > 59:
        return None
    return dt.time(h, mi)


def parse_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def fmt_time(t: dt.time | None, ampm: bool = True) -> str:
    if t is None:
        return ""
    h = t.hour % 12 or 12
    s = f"{h}:{t.minute:02d}"
    return f"{s} {'PM' if t.hour >= 12 else 'AM'}" if ampm else s


def fmt_range(a: dt.time | None, b: dt.time | None) -> str:
    if a and b:
        same = (a.hour >= 12) == (b.hour >= 12)
        return f"{fmt_time(a, not same)} – {fmt_time(b)}"
    return fmt_time(a or b)


def short_time(t: dt.time | None) -> str:
    if t is None:
        return ""
    h = t.hour % 12 or 12
    return f"{h}:{t.minute:02d}{'p' if t.hour >= 12 else 'a'}"


def tstr(t: dt.time | None) -> str:
    return t.strftime("%H:%M") if t else ""


def minutes(a: dt.time | None, b: dt.time | None) -> int | None:
    if not (a and b):
        return None
    d = (b.hour * 60 + b.minute) - (a.hour * 60 + a.minute)
    return d if d >= 0 else None


def day_label(d: dt.date, long: bool = True) -> str:
    return d.strftime("%A, %B ") + str(d.day) if long else d.strftime("%a ") + str(d.day)


# --------------------------------------------------------------------------- colors

def theme_colors(wb) -> list[str]:
    """Theme palette in Excel's index order (lt1, dk1, lt2, dk2, accent1-6, hlink, folHlink)."""
    fallback = ["FFFFFF", "000000", "E8E8E8", "0E2841", "156082", "E97132", "196B24", "0F9ED5", "A02B93", "4EA72E", "467886", "96607D"]
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return fallback
    try:
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = ET.fromstring(raw).find(".//a:clrScheme", ns)
        vals = {}
        for child in scheme:
            tag = child.tag.split("}")[1]
            clr = child[0]
            vals[tag] = clr.get("val") if clr.tag.endswith("srgbClr") else clr.get("lastClr", "000000")
        order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]
        return [vals.get(k, fallback[i]) for i, k in enumerate(order)]
    except Exception:
        return fallback


def apply_tint(hex6: str, tint: float) -> str:
    if not tint:
        return hex6
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "".join(f"{round(c * 255):02X}" for c in (r, g, b))


def cell_fill(cell, theme: list[str]) -> str | None:
    """Resolved solid fill as RRGGBB, or None for no/white fill."""
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return None
    c = fill.fgColor
    rgb = None
    try:
        if c.type == "rgb" and isinstance(c.rgb, str):
            rgb = c.rgb[-6:]
        elif c.type == "theme" and c.theme is not None and c.theme < len(theme):
            rgb = apply_tint(theme[c.theme], c.tint or 0)
        elif c.type == "indexed" and c.indexed is not None and c.indexed < len(COLOR_INDEX):
            rgb = COLOR_INDEX[c.indexed][-6:]
    except Exception:
        rgb = None
    if not rgb or (c.type == "rgb" and c.rgb == "00000000") or rgb.upper() == "FFFFFF":
        return None
    return rgb.upper()


def fill_kind(rgb: str | None) -> str:
    if not rgb:
        return ""
    r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    if s < 0.12 or max(r, g, b) - min(r, g, b) < 0.08:
        return "blocked" if l < 0.93 else ""
    if 0.11 <= h <= 0.19 and l > 0.45:
        return "hl"
    return "custom"


# --------------------------------------------------------------------------- workbook -> model

def header_map(ws, row: int = 1) -> dict[str, int]:
    out = {}
    for c in ws[row]:
        name = clean(c.value).lower()
        if name and name not in out:
            out[name] = c.column
    return out


def classify(ws) -> str:
    h = header_map(ws)
    keys = list(h)
    if any("session" in k for k in keys) and "date" in h and "start" in h:
        return "schedule"
    if keys[:4] == ["day", "date", "start", "end"]:
        return "grid"
    return "table"


def find_col(h: dict[str, int], *names: str) -> int | None:
    for n in names:
        if n in h:
            return h[n]
    for n in names:
        for k, v in h.items():
            if n in k:
                return v
    return None


def split_people(s: str) -> list[str]:
    return [p.strip() for p in re.split(r"\s*(?:\+|,|&|/|\band\b)\s*", s) if p.strip()]


def split_audience(s: str) -> list[str]:
    return [p.strip() for p in re.split(r"\s*[,;\n]\s*", s) if p.strip()]


TYPE_PATTERNS = (
    (re.compile(r"^(General Session)\s*\((.+)\)\s*$", re.S), None),
    (re.compile(r"^(Market Meeting[^(]*?)\s*\((.+)\)\s*$", re.S), None),
    (re.compile(r"^(Workshop|CE|Early Bird|Breakout)\s*:\s*(.+?)\s*(?:\([^()]*\))?\s*$", re.S), None),
)


def split_title(title: str) -> tuple[str, str]:
    for pat, _ in TYPE_PATTERNS:
        m = pat.match(title)
        if m:
            return m.group(1).strip(), m.group(2).strip()
    return "", title


def read_schedule(ws, theme) -> dict:
    h = header_map(ws)
    col = {
        "forum": find_col(h, "forum"),
        "code": find_col(h, "code"),
        "title": find_col(h, "session name", "session"),
        "presenters": find_col(h, "presenters", "presenter"),
        "date": find_col(h, "date"),
        "start": find_col(h, "start"),
        "end": find_col(h, "end"),
        "location": find_col(h, "location", "room"),
        "producer": find_col(h, "room producer", "producer"),
        "audience": find_col(h, "audience"),
    }
    skip = {v for v in col.values() if v} | {find_col(h, "time")}
    extra_cols = [(name, c) for name, c in h.items() if c not in skip]
    sessions = []
    for r in range(2, ws.max_row + 1):
        def v(key):
            c = col.get(key)
            return ws.cell(r, c).value if c else None

        title = clean(v("title"))
        date = parse_date(v("date"))
        if not title and not date:
            continue
        start, end = parse_time(v("start")), parse_time(v("end"))
        code_cell = ws.cell(r, col["code"]) if col["code"] else None
        title_cell = ws.cell(r, col["title"]) if col["title"] else None
        kind, name = split_title(title)
        producer = clean(v("producer"))
        audience = clean(v("audience"))
        sessions.append({
            "row": r,
            "forum": clean(v("forum")),
            "code": clean(v("code")),
            "code_fill": cell_fill(code_cell, theme) if code_cell is not None else None,
            "title": title,
            "title_fill": cell_fill(title_cell, theme) if title_cell is not None else None,
            "struck": bool(title_cell is not None and title_cell.font and title_cell.font.strike),
            "kind": kind,
            "name": name,
            "presenters": clean(v("presenters")),
            "date": date.isoformat() if date else "",
            "start": tstr(start),
            "end": tstr(end),
            "location": clean(v("location")),
            "producer": producer,
            "producers": split_people(producer),
            "audience": audience,
            "audiences": split_audience(audience),
            "extra": {name.title(): clean(ws.cell(r, c).value) for name, c in extra_cols if clean(ws.cell(r, c).value)},
        })
    sessions.sort(key=lambda s: (s["date"] or "9999", s["start"] or "99", s["row"]))
    return {"type": "schedule", "sessions": sessions}


def read_grid(ws, theme) -> dict:
    headers = [clean(c.value) for c in ws[1]]
    last = max(i for i, hdr in enumerate(headers) if hdr) + 1
    room_cols = list(range(5, last + 1))
    rooms = [headers[c - 1] for c in room_cols]

    merges = {}
    covered = set()
    for rng in ws.merged_cells.ranges:
        merges[(rng.min_row, rng.min_col)] = rng
        for rr in range(rng.min_row, rng.max_row + 1):
            for cc in range(rng.min_col, rng.max_col + 1):
                if (rr, cc) != (rng.min_row, rng.min_col):
                    covered.add((rr, cc))

    # Pass 1: sort rows into day groups, sub-headers and notes.
    days, notes, pending_sub = [], [], None
    current = None
    for r in range(2, ws.max_row + 1):
        date = parse_date(ws.cell(r, 2).value)
        room_text = [clean(ws.cell(r, c).value) for c in room_cols]
        first = clean(ws.cell(r, 1).value)
        if date:
            if current is None or current["date"] != date.isoformat():
                current = {"date": date.isoformat(), "rows": [], "sub": pending_sub}
                days.append(current)
                pending_sub = None
            current["rows"].append(r)
            continue
        if any(room_text):
            pending_sub = r
            current = None
            continue
        current = None
        if first and first.lower().strip() not in WEEKDAYS:
            fill = cell_fill(ws.cell(r, 1), theme)
            notes.append({"text": first, "fill": fill, "kind": fill_kind(fill), "bold": bool(ws.cell(r, 1).font.b)})

    def cell_model(r, c, rows_in_day):
        rng = merges.get((r, c))
        rowspan = colspan = 1
        if rng:
            colspan = min(rng.max_col, room_cols[-1]) - c + 1
            rowspan = sum(1 for rr in rows_in_day if rng.min_row <= rr <= rng.max_row)
        cell = ws.cell(r, c)
        fill = cell_fill(cell, theme)
        return {
            "text": clean(cell.value),
            "fill": fill,
            "kind": fill_kind(fill),
            "rowspan": max(rowspan, 1),
            "colspan": max(colspan, 1),
            "bold": bool(cell.font and cell.font.b),
            "struck": bool(cell.font and cell.font.strike),
        }

    out_days = []
    for d in days:
        rows_in_day = d["rows"]
        row_set = set(rows_in_day)
        sub = None
        if d["sub"]:
            sub = []
            for c in room_cols:
                if (d["sub"], c) in covered:
                    continue
                m = cell_model(d["sub"], c, [d["sub"]])
                sub.append({"text": m["text"], "colspan": m["colspan"]})
        out_rows = []
        for r in rows_in_day:
            cells = []
            for i, c in enumerate(room_cols):
                if (r, c) in covered:
                    # Merges that start above this day (or on a separator row) get re-anchored here.
                    origin = next((rng for rng in ws.merged_cells.ranges
                                   if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col), None)
                    first_row_in_day = origin and min(rr for rr in rows_in_day if origin.min_row <= rr <= origin.max_row)
                    if origin and origin.min_row not in row_set and first_row_in_day == r and origin.min_col == c:
                        m = cell_model(origin.min_row, c, rows_in_day)
                        m["rowspan"] = sum(1 for rr in rows_in_day if origin.min_row <= rr <= origin.max_row)
                        cells.append({**m, "col": i})
                    continue
                cells.append({**cell_model(r, c, rows_in_day), "col": i})
            out_rows.append({
                "row": r,
                "start": tstr(parse_time(ws.cell(r, 3).value)),
                "end": tstr(parse_time(ws.cell(r, 4).value)),
                "cells": cells,
            })
        out_days.append({"date": d["date"], "sub": sub, "rows": out_rows})
    return {"type": "grid", "rooms": rooms, "days": out_days, "notes": notes}


def read_table(ws, theme) -> dict:
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        vals = [clean(c.value.strftime("%a %b %d") if isinstance(c.value, dt.datetime) else c.value) for c in row]
        if any(vals):
            rows.append(vals)
    width = max((len(r) for r in rows), default=0)
    while width and not any(len(r) >= width and r[width - 1] for r in rows):
        width -= 1
    return {"type": "table", "rows": [r[:width] for r in rows]}


def load_model(path: str | Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    theme = theme_colors(wb)
    sheets = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        kind = classify(ws)
        reader = {"schedule": read_schedule, "grid": read_grid}.get(kind, read_table)
        data = reader(ws, theme)
        data["name"] = ws.title
        data["id"] = "sh-" + slug(ws.title)
        sheets.append(data)
    return {"sheets": sheets}


# --------------------------------------------------------------------------- change tracking

def friendly_name(sheet_name: str) -> str:
    return re.sub(r"^\d{4}\s+", "", sheet_name.replace("Sign Up", "Rehearsals").replace("Presenter ", "").strip())


def _index(model: dict) -> dict[str, tuple[str, dict]]:
    """Flatten a model to {stable key: (human label, fields)} for diffing."""
    out = {}
    for sh in model.get("sheets", []):
        if sh["type"] == "schedule":
            seen = {}
            for s in sh["sessions"]:
                base = f"{sh['id']}|{s['date']}|{s['start']}|{s['location']}"
                n = seen[base] = seen.get(base, 0) + 1
                d = dt.date.fromisoformat(s["date"]) if s["date"] else None
                label = f"{friendly_name(sh['name'])} · {day_label(d, False) if d else '?'} · {fmt_time(parse_time(s['start']))} · {s['location']}"
                out[f"{base}|{n}"] = (label, {
                    "Session": s["title"], "Code": s["code"], "Presenters": s["presenters"],
                    "Ends": fmt_time(parse_time(s["end"])), "Producer": s["producer"], "Audience": s["audience"],
                    "Highlight": "yes" if s["code_fill"] else "",
                })
        elif sh["type"] == "grid":
            for d in sh["days"]:
                sub = _room_names(sh, d)
                for row in d["rows"]:
                    for c in row["cells"]:
                        if not c["text"] and c["kind"] != "hl":
                            continue
                        room = sub[c["col"]]
                        label = f"{friendly_name(sh['name'])} · {day_label(dt.date.fromisoformat(d['date']), False)} · {fmt_time(parse_time(row['start']))} · {room}"
                        out[f"{sh['id']}|{d['date']}|{row['start']}|{c['col']}"] = (label, {
                            "Booking": c["text"], "Highlight": "yes" if c["kind"] == "hl" else "",
                        })
        else:
            for i, r in enumerate(sh["rows"]):
                out[f"{sh['id']}|{i}"] = (f"{sh['name']} · row {i + 1}", {"Row": " | ".join(r)})
    return out


def diff_models(old: dict | None, new: dict) -> list[dict]:
    if not old:
        return []
    a, b = _index(old), _index(new)
    changes = []
    for k in b.keys() - a.keys():
        label, f = b[k]
        changes.append({"key": k, "type": "added", "label": label, "detail": f.get("Session") or f.get("Booking") or f.get("Row", "")})
    for k in a.keys() - b.keys():
        label, f = a[k]
        changes.append({"key": k, "type": "removed", "label": label, "detail": f.get("Session") or f.get("Booking") or f.get("Row", "")})
    for k in a.keys() & b.keys():
        label, fa = a[k]
        fb = b[k][1]
        for field in fb:
            if fa.get(field, "") != fb.get(field, ""):
                changes.append({"key": k, "type": "changed", "label": label, "field": field,
                                "old": fa.get(field, ""), "new": fb.get(field, "")})
    changes.sort(key=lambda c: c["key"])
    return changes


# --------------------------------------------------------------------------- rendering

def e(s) -> str:
    return html.escape(str(s or ""), quote=True)


def multiline(s: str) -> str:
    return "<br>".join(e(line) for line in s.split("\n"))


def _room_names(sheet: dict, day: dict) -> list[str]:
    """Room label per column for a grid day, honouring a sub-header row if present."""
    names = list(sheet["rooms"])
    if day.get("sub"):
        i = 0
        for s in day["sub"]:
            for k in range(s["colspan"]):
                if i + k < len(names) and s["text"]:
                    names[i + k] = s["text"]
            i += s["colspan"]
    return names


def room_palette(model: dict) -> dict[str, str]:
    names = []
    for sh in model["sheets"]:
        if sh["type"] == "schedule":
            names += [s["location"] for s in sh["sessions"]]
        elif sh["type"] == "grid":
            names += sh["rooms"]
    pal, i = {}, 0
    for n in names:
        key = room_key(n)
        if key and key not in pal:
            pal[key] = ROOM_COLORS[i % len(ROOM_COLORS)]
            i += 1
    return pal


def room_key(name: str) -> str:
    n = name.lower()
    n = re.sub(r"\(.*?\)", "", n)
    for word in ("platinum", "south", "center", "centre", "north", "east", "west"):
        if word in n:
            return word
    return slug(n)


def room_color(pal: dict, name: str) -> str:
    return pal.get(room_key(name), "var(--muted)")


class Page:
    def __init__(self, model: dict, meta: dict):
        self.model = model
        self.meta = meta
        self.pal = room_palette(model)
        self.radios: list[str] = []
        self.css: list[str] = []
        self.changed_keys = {c["key"] for c in meta.get("recent_changes", [])}
        self.forum_audiences: dict[str, list[str]] = {}

    def radio(self, name: str, rid: str, checked: bool = False) -> None:
        self.radios.append(f'<input type="radio" class="ctl" name="{e(name)}" id="{e(rid)}"{" checked" if checked else ""}>')

    # ---- schedule
    def schedule(self, sh: dict) -> tuple[str, str]:
        sid = sh["id"]
        sessions = sh["sessions"]
        self.forum_audiences = {}
        for s in sessions:
            fa = self.forum_audiences.setdefault(s["forum"], [])
            fa += [a for a in s["audiences"] if a not in fa]
        dates = sorted({s["date"] for s in sessions if s["date"]})
        auds = []
        for s in sessions:
            for a in s["audiences"]:
                if a not in auds:
                    auds.append(a)
        prods = []
        for s in sessions:
            for p in s["producers"]:
                if p not in prods:
                    prods.append(p)

        # Radios: day, then audience, then producer (order matters for the ~ combinator rules).
        self.radio(f"{sid}-day", f"{sid}-d-all", True)
        for d in dates:
            self.radio(f"{sid}-day", f"{sid}-d-{d}")
        self.radio(f"{sid}-aud", f"{sid}-a-all", True)
        for a in auds:
            self.radio(f"{sid}-aud", f"{sid}-a-{slug(a)}")
        self.radio(f"{sid}-prod", f"{sid}-p-all", True)
        for p in prods:
            self.radio(f"{sid}-prod", f"{sid}-p-{slug(p)}")

        css = self.css
        for d in dates:
            css.append(f"#{sid}-d-{d}:checked~main #{sid}-view .sday:not(#{sid}-day-{d}){{display:none}}")
        for a in auds:
            css.append(f"#{sid}-a-{slug(a)}:checked~main #{sid}-view .card:not(.a-{slug(a)}){{display:none}}")
        for p in prods:
            css.append(f"#{sid}-p-{slug(p)}:checked~main #{sid}-view .card:not(.p-{slug(p)}){{display:none}}")
        # Hide empty time slots / show empty-day notes for every non-default filter combination.
        combos = [(a, p) for a in ["all"] + [slug(x) for x in auds] for p in ["all"] + [slug(x) for x in prods] if (a, p) != ("all", "all")]
        for a, p in combos:
            sel = f"#{sid}-a-{a}:checked~#{sid}-p-{p}:checked~main #{sid}-view"
            css.append(f"{sel} .slot:not(.h-{a}--{p}){{display:none}}")
            css.append(f"{sel} .sday:not(.h-{a}--{p}) .empty{{display:block}}")
            css.append(f"#{sid}-d-all:checked~#{sid}-a-{a}:checked~#{sid}-p-{p}:checked~main #{sid}-view .sday:not(.h-{a}--{p}){{display:none}}")
        for rid in [f"{sid}-a-all"] + [f"{sid}-a-{slug(a)}" for a in auds] + [f"{sid}-p-all"] + [f"{sid}-p-{slug(p)}" for p in prods] + [f"{sid}-d-all"] + [f"{sid}-d-{d}" for d in dates]:
            css.append(f"#{rid}:checked~header label[for={rid}],#{rid}:checked~main label[for={rid}]{{background:var(--chip-on);color:var(--chip-on-fg);border-color:var(--chip-on)}}")
        css.append(f"#{sid}-a-all:not(:checked)~header .fdot-{sid},#{sid}-p-all:not(:checked)~header .fdot-{sid}{{display:inline-block}}")

        # Toolbar (sticky): day chips + filter panel.
        tb = [f'<div class="tb-row chips scroll" role="radiogroup" aria-label="Day">',
              f'<label for="{sid}-d-all" class="chip">All days</label>']
        for d in dates:
            dd = dt.date.fromisoformat(d)
            tb.append(f'<label for="{sid}-d-{d}" class="chip" data-date="{d}">{e(dd.strftime("%a"))} <b>{dd.day}</b></label>')
        tb.append("</div>")
        tb.append(f'<details class="filters"><summary><span>Filter</span><i class="fdot fdot-{sid}"></i></summary><div class="fpanel">')
        tb.append('<div class="fgroup"><div class="flabel">Audience</div><div class="chips wrap">')
        tb.append(f'<label for="{sid}-a-all" class="chip">Everyone</label>')
        for a in auds:
            tb.append(f'<label for="{sid}-a-{slug(a)}" class="chip">{e(a)}</label>')
        tb.append('</div></div><div class="fgroup"><div class="flabel">Room producer</div><div class="chips wrap">')
        tb.append(f'<label for="{sid}-p-all" class="chip">All producers</label>')
        for p in prods:
            rooms = sorted({s["location"] for s in sessions if p in s["producers"]})
            dot = "".join(f'<i class="rdot" style="background:{room_color(self.pal, r)}"></i>' for r in rooms[:2])
            tb.append(f'<label for="{sid}-p-{slug(p)}" class="chip">{dot}{e(p)}</label>')
        tb.append("</div></div></div></details>")
        toolbar = "".join(tb)

        out = []
        for d in dates:
            day_sessions = [s for s in sessions if s["date"] == d]
            slots: list[list[dict]] = []
            for s in day_sessions:
                if slots and slots[-1][0]["start"] == s["start"]:
                    slots[-1].append(s)
                else:
                    slots.append([s])
            day_has = set()
            body = []
            for slot in slots:
                has = set()
                for s in slot:
                    for a in ["all"] + [slug(x) for x in s["audiences"]]:
                        for p in ["all"] + [slug(x) for x in s["producers"]]:
                            has.add(f"h-{a}--{p}")
                day_has |= has
                st = parse_time(slot[0]["start"])
                body.append(f'<div class="slot {" ".join(sorted(has))}"><div class="slot-time">{e(fmt_time(st))}</div><div class="slot-cards">')
                body += [self.card(sh, s) for s in slot]
                body.append("</div></div>")
            dd = dt.date.fromisoformat(d)
            out.append(f'<section class="sday {" ".join(sorted(day_has))}" id="{sid}-day-{d}" data-date="{d}">'
                       f'<h2 class="dayh"><span>{e(dd.strftime("%A"))}</span> {e(dd.strftime("%B"))} {dd.day}</h2>'
                       f'<p class="empty">Nothing on this day for the selected filter.</p>{"".join(body)}</section>')
        undated = [s for s in sessions if not s["date"]]
        if undated:
            out.append('<section class="sday h-all--all"><h2 class="dayh">No date</h2><div class="slot h-all--all"><div class="slot-cards">'
                       + "".join(self.card(sh, s) for s in undated) + "</div></div></section>")
        return toolbar, "".join(out)

    def card(self, sh: dict, s: dict) -> str:
        cls = ["card", f"f-{slug(s['forum'])}"] + [f"a-{slug(a)}" for a in s["audiences"]] + [f"p-{slug(p)}" for p in s["producers"]]
        key_base = f"{sh['id']}|{s['date']}|{s['start']}|{s['location']}"
        changed = any(k.startswith(key_base + "|") for k in self.changed_keys)
        if changed:
            cls.append("changed")
        if s["struck"]:
            cls.append("struck")
        st, en = parse_time(s["start"]), parse_time(s["end"])
        mins = minutes(st, en)
        code = ""
        if s["code"]:
            hl = fill_kind(s["code_fill"]) == "hl"
            code = f'<span class="code{" hl" if hl else ""}"{"" if hl or not s["code_fill"] else f" style=background:#{s['code_fill']}"}>{e(s["code"])}</span>'
        everyone = self.forum_audiences.get(s["forum"], [])
        if len(everyone) > 1 and sorted(s["audiences"]) == sorted(everyone):
            aud = f'<span class="aud all" title="{e(", ".join(everyone))}">All {e(s["forum"])} groups</span>'
        else:
            aud = "".join(f'<span class="aud">{e(a)}</span>' for a in s["audiences"])
        extra = "".join(f'<div class="xtra"><b>{e(k)}:</b> {multiline(v)}</div>' for k, v in s["extra"].items())
        color = room_color(self.pal, s["location"])
        search = " ".join([s["code"], s["title"], s["presenters"], s["location"], s["producer"], s["audience"], s["forum"]]).lower()
        return (
            f'<article class="{" ".join(cls)}" style="--rc:{color}" data-s="{s["date"]}T{s["start"]}" data-e="{s["date"]}T{s["end"]}" data-q="{e(search)}">'
            f'<div class="c-top"><span class="c-time">{e(fmt_range(st, en))}</span>'
            f'{f"<span class=dur>{mins} min</span>" if mins else ""}'
            f'<span class="forum">{e(s["forum"])}</span>{"<span class=upd>Updated</span>" if changed else ""}</div>'
            f'{f"<div class=kind>{e(s["kind"])}</div>" if s["kind"] else ""}'
            f'<h3 class="c-title">{code}{multiline(s["name"])}</h3>'
            f'{f"<div class=who>{multiline(s["presenters"])}</div>" if s["presenters"] else ""}'
            f'<div class="where"><span class="room"><i class="rdot" style="background:{color}"></i>{e(s["location"])}</span>'
            f'{f"<span class=prod>{e(s["producer"])}</span>" if s["producer"] else ""}</div>'
            f'{f"<div class=auds>{aud}</div>" if aud else ""}{extra}</article>'
        )

    # ---- rehearsal grids
    def grid(self, sh: dict) -> tuple[str, str]:
        sid = sh["id"]
        self.radio(f"{sid}-mode", f"{sid}-m-auto", True)
        self.radio(f"{sid}-mode", f"{sid}-m-list")
        self.radio(f"{sid}-mode", f"{sid}-m-grid")
        css = self.css
        v = f"~main #{sid}-view"
        css.append(f"#{sid}-m-list:checked{v} .rgrid-wrap,#{sid}-m-grid:checked{v} .rlist{{display:none}}")
        css.append(f"@media (min-width:700px){{#{sid}-m-auto:checked{v} .rlist{{display:none}}#{sid}-m-auto:checked~header label[for={sid}-m-grid]{{background:var(--seg-on);color:var(--seg-on-fg)}}}}")
        css.append(f"@media (max-width:699.98px){{#{sid}-m-auto:checked{v} .rgrid-wrap{{display:none}}#{sid}-m-auto:checked~header label[for={sid}-m-list]{{background:var(--seg-on);color:var(--seg-on-fg)}}}}")
        for m in ("list", "grid"):
            css.append(f"#{sid}-m-{m}:checked~header label[for={sid}-m-{m}]{{background:var(--seg-on);color:var(--seg-on-fg)}}")

        days = sh["days"]
        tb = ['<div class="tb-row chips scroll">']
        for d in days:
            dd = dt.date.fromisoformat(d["date"])
            tb.append(f'<a class="chip" href="#{sid}-day-{d["date"]}" data-date="{d["date"]}">{e(dd.strftime("%a"))} <b>{dd.day}</b></a>')
        tb.append(f'</div><div class="seg" role="radiogroup" aria-label="Layout"><label for="{sid}-m-list">List</label><label for="{sid}-m-grid">Grid</label></div>')
        toolbar = "".join(tb)

        out = []
        for d in days:
            dd = dt.date.fromisoformat(d["date"])
            names = _room_names(sh, d)
            out.append(f'<section class="rday" id="{sid}-day-{d["date"]}" data-date="{d["date"]}">'
                       f'<h2 class="dayh"><span>{e(dd.strftime("%A"))}</span> {e(dd.strftime("%B"))} {dd.day}</h2>')
            out.append(self.grid_table(sh, d))
            out.append(self.grid_list(sh, d, names))
            out.append("</section>")
        if sh["notes"]:
            out.append('<div class="notes">')
            for n in sh["notes"]:
                sw = f'<i class="sw {n["kind"]}"{f" style=background:#{n["fill"]}" if n["kind"] == "custom" else ""}></i>' if n["fill"] else ""
                out.append(f'<p class="note{" b" if n["bold"] else ""}">{sw}{multiline(n["text"])}</p>')
            out.append("</div>")
        return toolbar, "".join(out)

    def grid_cell_class(self, sh, d, row, c) -> tuple[str, str]:
        cls = [c["kind"]] if c["kind"] else []
        style = ""
        if c["kind"] == "custom":
            style = f' style="background:#{c["fill"]}"'
        if c["text"]:
            cls.append("booked")
        if c["bold"]:
            cls.append("b")
        if c["struck"]:
            cls.append("struck")
        if f"{sh['id']}|{d['date']}|{row['start']}|{c['col']}" in self.changed_keys:
            cls.append("changed")
        return " ".join(cls), style

    def grid_table(self, sh: dict, d: dict) -> str:
        rooms = sh["rooms"]
        p = ['<div class="rgrid-wrap"><table class="rgrid"><thead><tr><th class="t">Time</th>']
        for r in rooms:
            p.append(f'<th style="--rc:{room_color(self.pal, r)}">{e(r)}</th>')
        p.append("</tr>")
        if d.get("sub"):
            p.append('<tr class="sub"><th class="t"></th>')
            for s in d["sub"]:
                p.append(f'<th colspan="{s["colspan"]}" style="--rc:{room_color(self.pal, s["text"])}">{e(s["text"])}</th>')
            p.append("</tr>")
        p.append("</thead><tbody>")
        by_row = {row["row"]: row for row in d["rows"]}
        rows = d["rows"]
        for i, row in enumerate(rows):
            st, en = parse_time(row["start"]), parse_time(row["end"])
            p.append(f'<tr data-s="{d["date"]}T{row["start"]}" data-e="{d["date"]}T{row["end"]}"><th class="t">{e(short_time(st))}<small>{e(short_time(en))}</small></th>')
            for c in row["cells"]:
                cls, style = self.grid_cell_class(sh, d, row, c)
                span = (f' rowspan="{c["rowspan"]}"' if c["rowspan"] > 1 else "") + (f' colspan="{c["colspan"]}"' if c["colspan"] > 1 else "")
                if c["rowspan"] > 1:
                    last = rows[min(i + c["rowspan"] - 1, len(rows) - 1)]
                    cls += " multi"
                    end_attr = f' data-s="{d["date"]}T{row["start"]}" data-e="{d["date"]}T{last["end"]}"'
                else:
                    end_attr = ""
                p.append(f'<td class="{cls}"{style}{span}{end_attr} data-q="{e(c["text"].lower())}">{multiline(c["text"])}</td>')
            p.append("</tr>")
        p.append("</tbody></table></div>")
        return "".join(p)

    def grid_list(self, sh: dict, d: dict, names: list[str]) -> str:
        rows = d["rows"]
        entries = []
        for i, row in enumerate(rows):
            for c in row["cells"]:
                if not c["text"]:
                    continue
                last = rows[min(i + c["rowspan"] - 1, len(rows) - 1)]
                room = names[c["col"]]
                if c["colspan"] > 1:
                    joined = [names[c["col"] + k] for k in range(c["colspan"]) if c["col"] + k < len(names)]
                    room = room if len(set(joined)) == 1 else " + ".join(dict.fromkeys(joined))
                entries.append((row["start"], c["col"], row, last, c, room))
        entries.sort(key=lambda x: (x[0], x[1]))
        if not entries:
            return '<div class="rlist"><p class="empty show">No rehearsals booked.</p></div>'
        p = ['<div class="rlist">']
        for start, col, row, last, c, room in entries:
            st, en = parse_time(row["start"]), parse_time(last["end"])
            cls, style = self.grid_cell_class(sh, d, row, c)
            p.append(f'<div class="rentry {cls}" style="--rc:{room_color(self.pal, room)}" data-s="{d["date"]}T{row["start"]}" data-e="{d["date"]}T{last["end"]}" data-q="{e((c["text"] + " " + room).lower())}">'
                     f'<div class="rt">{e(fmt_range(st, en))}</div><div class="rroom"><i class="rdot" style="background:{room_color(self.pal, room)}"></i>{e(room)}</div>'
                     f'<div class="rtext">{multiline(c["text"])}{"<span class=upd>Updated</span>" if "changed" in cls.split() else ""}</div></div>')
        p.append("</div>")
        return "".join(p)

    # ---- generic table fallback
    def table(self, sh: dict) -> tuple[str, str]:
        rows = sh["rows"]
        if not rows:
            return "", '<p class="empty show">This sheet is empty.</p>'
        p = ['<div class="rgrid-wrap"><table class="rgrid plain"><thead><tr>']
        p += [f"<th>{e(h)}</th>" for h in rows[0]]
        p.append("</tr></thead><tbody>")
        for r in rows[1:]:
            p.append("<tr>" + "".join(f'<td data-q="{e(v.lower())}">{multiline(v)}</td>' for v in r) + "</tr>")
        p.append("</tbody></table></div>")
        return "", "".join(p)

    # ---- changes view
    def changes(self) -> str:
        batches = self.meta.get("change_batches", [])
        if not batches:
            return ('<section class="chg"><p class="empty show">No changes recorded yet. When someone edits the spreadsheet in Dropbox, '
                    'what changed shows up here and the affected sessions get an <span class="upd">Updated</span> tag for 24 hours.</p></section>')
        p = ['<section class="chg">']
        for b in batches:
            when = dt.datetime.fromisoformat(b["at"])
            p.append(f'<h2 class="dayh"><span>{e(when.strftime("%a %b"))} {when.day}</span> {e(fmt_time(when.time()))} · {len(b["changes"])} change{"s" if len(b["changes"]) != 1 else ""}</h2><ul class="clist">')
            for c in b["changes"][:200]:
                if c["type"] == "changed":
                    body = f'<b>{e(c["field"])}</b>: <del>{multiline(c["old"]) or "—"}</del> → <ins>{multiline(c["new"]) or "—"}</ins>'
                else:
                    body = multiline(c.get("detail") or "")
                p.append(f'<li class="ct-{c["type"]}"><span class="ctag">{e(c["type"].title())}</span><div><div class="clabel">{e(c["label"])}</div><div>{body}</div></div></li>')
            if len(b["changes"]) > 200:
                p.append(f'<li>…and {len(b["changes"]) - 200} more</li>')
            p.append("</ul>")
        p.append("</section>")
        return "".join(p)

    # ---- page
    def render(self) -> str:
        meta = self.meta
        views = []
        for sh in self.model["sheets"]:
            fn = {"schedule": self.schedule, "grid": self.grid}.get(sh["type"], self.table)
            label = friendly_name(sh["name"])
            views.append((sh["id"], label, *fn(sh)))
        n_recent = len(meta.get("recent_changes", []))
        views.append(("changes", "Changes", "", self.changes()))

        view_radios = []
        for i, (vid, *_rest) in enumerate(views):
            view_radios.append(f'<input type="radio" class="ctl" name="view" id="v-{vid}"{" checked" if i == 0 else ""}>')
            self.css.append(f"#v-{vid}:checked~main .view:not(#{vid}-view),#v-{vid}:checked~header .tb:not(.tb-{vid}){{display:none}}")
            self.css.append(f"#v-{vid}:checked~header label[for=v-{vid}]{{color:var(--fg);border-color:var(--accent)}}")

        tabs = "".join(
            f'<label for="v-{vid}" class="tab">{tab_label(label)}{f"<span class=badge>{n_recent}</span>" if vid == "changes" and n_recent else ""}</label>'
            for vid, label, *_ in views)
        toolbars = "".join(f'<div class="tb tb-{vid}">{tb}</div>' for vid, _l, tb, _b in views if tb)
        bodies = "".join(f'<div class="view" id="{vid}-view">{body}</div>' for vid, _l, _t, body in views)

        checked = dt.datetime.fromisoformat(meta["checked_at"])
        changed = dt.datetime.fromisoformat(meta["changed_at"]) if meta.get("changed_at") else None
        stamp = f"{checked.strftime('%a %b')} {checked.day}, {fmt_time(checked.time())}"
        err = meta.get("error")
        status = (f'<div class="alert">Couldn’t reach the Dropbox file since {e(meta.get("error_since_label", ""))}: {e(err)}. '
                  f'Showing the last good copy.</div>' if err else "")
        tz = meta.get("tz", "")
        title = meta.get("title", "Schedules + Rehearsals")
        dropbox = meta.get("dropbox_url", "")

        values = dict(
            title=e(title),
            eyebrow=e(meta.get("eyebrow", "")),
            radios="\n".join(view_radios + self.radios),
            css=CSS + "\n" + "\n".join(self.css),
            tabs=break_lines(tabs),
            toolbars=break_lines(toolbars),
            bodies=break_lines(bodies),
            stamp=e(stamp + (f" {tz}" if tz else "")),
            checked_iso=e(checked.astimezone().isoformat()),
            changed=e(f"{changed.strftime('%a %b')} {changed.day}, {fmt_time(changed.time())}" if changed else "—"),
            dropbox=e(dropbox),
            status=status,
            version=e(BUILD_VERSION),
            js=JS,
        )
        return re.sub(r"@@(\w+)@@", lambda m: values[m.group(1)], TEMPLATE)


def tab_label(label: str) -> str:
    """Full label on wide screens, abbreviated (Co./Fr.) on phones."""
    out = e(label)
    for word, short in (("Company", "Co."), ("Franchise", "Fr.")):
        out = out.replace(word, f'<span class="lg">{word}</span><span class="sm">{short}</span>')
    return out


BLOCK_TAG = re.compile(r"<(?=(?:article|section|div|details|table|thead|tbody|tr|ul|li|label|a|h2|p)[\s>])")


def break_lines(markup: str) -> str:
    """Start each block-level element on its own line.

    Browsers ignore this whitespace, but text editors and apps that open .html as
    code (e.g. Pythonista on iPad) can crash on a single 100 KB line.
    """
    return BLOCK_TAG.sub("\n<", markup)


def render(model: dict, meta: dict) -> str:
    return Page(model, meta).render()


# --------------------------------------------------------------------------- static assets

TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="color-scheme" content="light dark">
<meta name="format-detection" content="telephone=no">
<meta name="generator" content="HRB schedule viewer @@version@@">
<title>@@title@@</title>
<style>@@css@@</style>
</head>
<body id="top">
@@radios@@
<div class="brand">
    <div class="eyebrow">@@eyebrow@@</div>
    <h1>@@title@@</h1>
    <div class="sync"><span class="live"></span><span>Synced from Dropbox <b id="checked" data-iso="@@checked_iso@@">@@stamp@@</b></span>
      <span class="sep">·</span><span>last edit seen @@changed@@</span>
      <a class="open" href="@@dropbox@@">Open original ↗</a></div>
  @@status@@
  <div class="stale" id="stale" hidden></div>
</div>
<header class="top">
  <nav class="tabs scroll">@@tabs@@</nav>
  <div class="toolbars">@@toolbars@@<div class="searchbar" id="searchbar" hidden><input type="search" id="q" placeholder="Search names, sessions, rooms…" autocomplete="off" autocorrect="off" spellcheck="false"><span id="qcount"></span><button type="button" id="nowbtn" hidden>Now</button></div></div>
</header>
<main>@@bodies@@</main>
<footer>
  <p><i class="sw hl"></i>Highlighted in the sheet <i class="sw blocked"></i>Not available <i class="sw changed"></i>Changed in the last 24 h</p>
  <p>Built from the producers spreadsheet in Dropbox by the Mac mini, which checks it every 2 minutes and rewrites this file when it changes. Viewer @@version@@.</p>
  <p><a href="#top">Back to top ↑</a></p>
</footer>
<script>@@js@@</script>
</body>
</html>
"""


CSS = r"""
:root{
  --bg:#f5f6f8;--card:#fff;--fg:#0f172a;--muted:#5b6475;--line:#e3e6ec;--accent:#002060;--accent-2:#1d4ed8;
  --chip:#fff;--chip-fg:#334155;--chip-on:#002060;--chip-on-fg:#fff;--seg-on:#002060;--seg-on-fg:#fff;
  --hl:#fff36b;--hl-fg:#1a1a00;--blocked:#d9dce2;--blocked-fg:#4b5563;--chg:#f97316;--now:#16a34a;
  --company:#1d4ed8;--franchise:#15803d;--hdr:rgba(245,246,248,.92);--shadow:0 1px 2px rgba(15,23,42,.06),0 1px 1px rgba(15,23,42,.04);
  color-scheme:light dark;
}
@media (prefers-color-scheme:dark){:root{
  --bg:#0b0f17;--card:#151b26;--fg:#e8ecf3;--muted:#98a2b3;--line:#263042;--accent:#8fb3ff;--accent-2:#8fb3ff;
  --chip:#151b26;--chip-fg:#cbd5e1;--chip-on:#8fb3ff;--chip-on-fg:#08101f;--seg-on:#8fb3ff;--seg-on-fg:#08101f;
  --hl:#5c5300;--hl-fg:#fff7a8;--blocked:#232a36;--blocked-fg:#8b95a7;--chg:#fb923c;--now:#22c55e;
  --company:#8fb3ff;--franchise:#4ade80;--hdr:rgba(11,15,23,.92);--shadow:none;
}}
*{box-sizing:border-box}
html{-webkit-text-size-adjust:100%;text-size-adjust:100%;scroll-padding-top:150px}
body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.4 -apple-system,BlinkMacSystemFont,"SF Pro Text","Segoe UI",Roboto,Helvetica,Arial,sans-serif;
  padding:0 max(12px,env(safe-area-inset-right)) env(safe-area-inset-bottom) max(12px,env(safe-area-inset-left))}
a{color:var(--accent-2)}
.ctl{position:fixed;top:0;left:0;opacity:0;pointer-events:none;width:1px;height:1px;margin:0}
[hidden]{display:none!important}
.scroll{overflow-x:auto;-webkit-overflow-scrolling:touch;scrollbar-width:none}
.scroll::-webkit-scrollbar{display:none}

.brand{padding:calc(env(safe-area-inset-top) + 14px) 0 2px}
header.top{position:sticky;top:0;z-index:20;margin:0 -12px;padding:env(safe-area-inset-top) 12px 0;background:var(--hdr);
  -webkit-backdrop-filter:saturate(1.4) blur(14px);backdrop-filter:saturate(1.4) blur(14px);border-bottom:1px solid var(--line)}
.brand{max-width:1200px;margin:0 auto}
.eyebrow{color:#2f9e44;font-weight:600;font-size:12px;letter-spacing:.02em}
h1{margin:0;font-size:20px;line-height:1.2;color:var(--accent);letter-spacing:-.01em}
.sync{display:flex;flex-wrap:wrap;align-items:center;gap:2px 6px;color:var(--muted);font-size:12px;margin-top:3px}
.sync b{font-weight:600;color:var(--fg)}
.sync .open{margin-left:auto;text-decoration:none;font-weight:600}
.live{width:7px;height:7px;border-radius:50%;background:var(--now);display:inline-block}
.alert,.stale{max-width:1200px;margin:8px auto 0;padding:8px 10px;border-radius:10px;background:#fff1e6;color:#7c2d12;font-size:13px;border:1px solid #fed7aa}
@media (prefers-color-scheme:dark){.alert,.stale{background:#3a1d0b;color:#fed7aa;border-color:#7c2d12}}
.tabs{display:flex;gap:18px;max-width:1200px;margin:2px auto 0;white-space:nowrap}
.tab .sm{display:none}
.tab{padding:8px 0 9px;font-weight:600;font-size:14px;color:var(--muted);border-bottom:2.5px solid transparent;cursor:pointer;-webkit-user-select:none;user-select:none}
.badge{display:inline-block;min-width:18px;padding:0 5px;margin-left:5px;border-radius:9px;background:var(--chg);color:#fff;font-size:11px;line-height:18px;text-align:center}
.toolbars{max-width:1200px;margin:0 auto}
.tb{display:flex;align-items:center;gap:8px;padding:7px 0}
.tb-row{flex:1 1 0;min-width:0}
.chips{display:flex;gap:6px}
.chips.scroll{flex-wrap:nowrap}
.chips.wrap{flex-wrap:wrap}
.chip{flex:0 0 auto;display:inline-flex;align-items:center;gap:4px;padding:6px 11px;border-radius:999px;border:1px solid var(--line);background:var(--chip);color:var(--chip-fg);
  font-size:13px;white-space:nowrap;cursor:pointer;text-decoration:none;-webkit-user-select:none;user-select:none;-webkit-tap-highlight-color:transparent}
.chip b{font-weight:700}
.chip.today::after,.chip.today::after{content:"";width:5px;height:5px;border-radius:50%;background:var(--now);margin-left:2px}
.seg{display:inline-flex;border:1px solid var(--line);border-radius:9px;overflow:hidden;background:var(--chip);flex:0 0 auto}
.seg label{padding:6px 12px;font-size:13px;font-weight:600;color:var(--chip-fg);cursor:pointer;-webkit-user-select:none;user-select:none}
details.filters{position:relative;flex:0 0 auto}
details.filters summary{list-style:none;display:inline-flex;align-items:center;gap:6px;padding:6px 12px;border:1px solid var(--line);border-radius:999px;background:var(--chip);
  font-size:13px;font-weight:600;color:var(--chip-fg);cursor:pointer;-webkit-user-select:none;user-select:none}
details.filters summary::-webkit-details-marker{display:none}
details.filters summary::after{content:"";width:6px;height:6px;border-right:1.5px solid currentColor;border-bottom:1.5px solid currentColor;transform:rotate(45deg) translateY(-2px)}
details.filters[open] summary::after{transform:rotate(-135deg) translateY(-1px)}
.fdot{display:none;width:7px;height:7px;border-radius:50%;background:var(--chg)}
.fpanel{position:absolute;right:0;top:calc(100% + 6px);width:min(560px,calc(100vw - 24px));max-height:65vh;overflow:auto;background:var(--card);border:1px solid var(--line);
  border-radius:14px;padding:12px;box-shadow:0 12px 32px rgba(0,0,0,.18);z-index:30}
.fgroup+.fgroup{margin-top:12px}
.flabel{font-size:11px;font-weight:700;text-transform:uppercase;letter-spacing:.06em;color:var(--muted);margin-bottom:6px}
.rdot{display:inline-block;width:8px;height:8px;border-radius:50%;margin-right:5px;flex:0 0 auto}
.chip .rdot{margin-right:0}
.searchbar{display:flex;align-items:center;gap:8px;padding:0 0 8px}
.searchbar input{flex:1;min-width:0;font:inherit;font-size:16px;padding:5px 12px;border-radius:10px;border:1px solid var(--line);background:var(--card);color:var(--fg)}
.searchbar span{font-size:12px;color:var(--muted);white-space:nowrap}
.searchbar button{font:inherit;font-size:13px;font-weight:700;padding:6px 12px;border-radius:999px;border:0;background:var(--now);color:#fff}

main{max-width:1200px;margin:0 auto;padding-bottom:24px}
.dayh{position:relative;margin:22px 0 10px;font-size:17px;font-weight:600;color:var(--muted);scroll-margin-top:170px}
.dayh span{color:var(--fg);font-weight:700}
.empty{display:none;color:var(--muted);font-style:italic;margin:6px 0}
.empty.show{display:block}

.slot{display:grid;grid-template-columns:1fr;gap:6px;margin-bottom:12px}
.slot-time{font-size:12px;font-weight:700;color:var(--muted);text-transform:uppercase;letter-spacing:.04em}
.slot-cards{display:grid;grid-template-columns:repeat(auto-fill,minmax(280px,1fr));gap:10px}
@media (min-width:700px){.slot{grid-template-columns:76px 1fr;gap:12px}.slot-time{padding-top:12px;text-align:right}}
.card{position:relative;background:var(--card);border:1px solid var(--line);border-left:4px solid var(--rc,var(--line));border-radius:12px;padding:10px 12px 11px;box-shadow:var(--shadow);min-width:0}
.card.changed{box-shadow:0 0 0 2px var(--chg)}
.card.now{box-shadow:0 0 0 2px var(--now)}
.card.struck .c-title{text-decoration:line-through;opacity:.6}
.c-top{display:flex;flex-wrap:wrap;align-items:center;gap:6px;font-size:13px}
.c-time{font-weight:700;font-variant-numeric:tabular-nums}
.dur{color:var(--muted);font-size:12px}
.forum{margin-left:auto;font-size:10.5px;font-weight:700;text-transform:uppercase;letter-spacing:.06em;color:var(--muted)}
.f-company .forum{color:var(--company)}.f-franchise .forum{color:var(--franchise)}
.upd,.nowtag{display:inline-block;font-size:10.5px;font-weight:700;text-transform:uppercase;letter-spacing:.05em;padding:1px 6px;border-radius:5px;background:var(--chg);color:#fff;margin-left:6px;vertical-align:1px}
.nowtag{background:var(--now)}
.kind{margin-top:6px;font-size:11px;font-weight:700;text-transform:uppercase;letter-spacing:.06em;color:var(--muted)}
.c-title{margin:2px 0 0;font-size:16px;line-height:1.3;font-weight:650}
.code{display:inline-block;font-size:12px;font-weight:700;padding:1px 6px;border-radius:5px;background:var(--bg);border:1px solid var(--line);margin-right:6px;vertical-align:2px;font-variant-numeric:tabular-nums}
.code.hl{background:var(--hl);color:var(--hl-fg);border-color:transparent}
.who{margin-top:5px;font-size:14px}
.where{display:flex;flex-wrap:wrap;gap:4px 12px;margin-top:6px;font-size:13px;color:var(--muted)}
.room{display:inline-flex;align-items:center;color:var(--fg);font-weight:600}
.prod::before{content:"Producer · ";opacity:.8}
.auds{display:flex;flex-wrap:wrap;gap:4px;margin-top:7px}
.aud{font-size:11.5px;padding:1px 7px;border-radius:999px;background:var(--bg);border:1px solid var(--line);color:var(--muted)}
.xtra{font-size:12.5px;margin-top:4px;color:var(--muted)}

.rgrid-wrap{overflow-x:auto;-webkit-overflow-scrolling:touch;border:1px solid var(--line);border-radius:12px;background:var(--card)}
.rgrid{border-collapse:separate;border-spacing:0;width:100%;min-width:560px;font-size:13px}
.rgrid th,.rgrid td{border-bottom:1px solid var(--line);border-right:1px solid var(--line);padding:6px 8px;vertical-align:top;text-align:left}
.rgrid tr>*:last-child{border-right:0}
.rgrid tbody tr:last-child>*{border-bottom:0}
.rgrid thead th{background:var(--accent);color:#fff;font-weight:700;font-size:12.5px;border-bottom:3px solid var(--rc,var(--accent));border-right-color:rgba(255,255,255,.18)}
@media (prefers-color-scheme:dark){.rgrid thead th{background:#1c2536;color:var(--fg)}}
.rgrid thead tr.sub th{background:var(--bg);color:var(--fg);font-size:12px}
.rgrid .t{position:sticky;left:0;z-index:1;width:62px;min-width:62px;background:var(--card);font-weight:700;font-variant-numeric:tabular-nums;white-space:nowrap}
.rgrid thead .t{background:var(--accent);z-index:2}
@media (prefers-color-scheme:dark){.rgrid thead .t{background:#1c2536}}
.rgrid thead tr.sub .t{background:var(--bg)}
.rgrid .t small{display:block;font-weight:500;color:var(--muted);font-size:11px}
.rgrid td{width:calc((100% - 62px)/4);min-width:110px}
.rgrid td.blocked,.rentry.blocked{background:repeating-linear-gradient(135deg,var(--blocked) 0 6px,transparent 6px 11px);color:var(--blocked-fg)}
.rgrid td.blocked.booked{background:var(--blocked)}
.rgrid td.hl,.rentry.hl{background:var(--hl);color:var(--hl-fg)}
.rgrid td.booked{font-weight:550}
.rgrid td.b{font-weight:700}
.rgrid td.struck,.rentry.struck .rtext{text-decoration:line-through}
.rgrid td.changed{box-shadow:inset 0 0 0 2px var(--chg)}
.rgrid tr.now>.t{color:var(--now)}
.rgrid tr.now>.t::before{content:"";position:absolute;left:0;top:0;bottom:0;width:3px;background:var(--now)}
.rgrid td.now{box-shadow:inset 0 0 0 2px var(--now)}
.rgrid.plain td{min-width:90px}
.match{outline:2px solid var(--accent-2);outline-offset:-2px}

.rlist{display:flex;flex-direction:column;gap:8px}
.rentry{display:grid;grid-template-columns:auto 1fr;gap:2px 10px;background:var(--card);border:1px solid var(--line);border-left:4px solid var(--rc,var(--line));border-radius:12px;padding:9px 12px;box-shadow:var(--shadow)}
.rentry .rt{font-weight:700;font-size:13px;font-variant-numeric:tabular-nums}
.rentry .rroom{font-size:12.5px;color:var(--muted);display:flex;align-items:center;justify-self:end;text-align:right}
.rentry .rtext{grid-column:1/-1;font-size:15px;font-weight:550}
.rentry.changed{box-shadow:0 0 0 2px var(--chg)}
.rentry.now{box-shadow:0 0 0 2px var(--now)}

.notes{margin-top:18px;padding:10px 12px;border-radius:12px;background:var(--card);border:1px solid var(--line)}
.note{margin:4px 0;font-size:13.5px;display:flex;align-items:flex-start;gap:8px}
.note.b{font-weight:700}
.sw{display:inline-block;width:14px;height:14px;border-radius:4px;flex:0 0 auto;vertical-align:-2px;margin:0 4px 0 10px;border:1px solid var(--line)}
.note .sw{margin:2px 0 0}
.sw.hl{background:var(--hl)}
.sw.blocked{background:repeating-linear-gradient(135deg,var(--blocked) 0 3px,transparent 3px 6px)}
.sw.changed{background:var(--chg)}
footer .sw:first-child{margin-left:0}

.chg .clist{list-style:none;margin:0;padding:0;display:flex;flex-direction:column;gap:8px}
.chg li{display:flex;gap:10px;background:var(--card);border:1px solid var(--line);border-radius:12px;padding:9px 12px;font-size:14px}
.chg .clabel{font-size:12.5px;color:var(--muted);margin-bottom:2px}
.ctag{flex:0 0 auto;align-self:flex-start;font-size:10.5px;font-weight:700;text-transform:uppercase;letter-spacing:.05em;padding:2px 6px;border-radius:5px;background:var(--bg);border:1px solid var(--line)}
.ct-added .ctag{background:#dcfce7;color:#14532d;border-color:transparent}
.ct-removed .ctag{background:#fee2e2;color:#7f1d1d;border-color:transparent}
.ct-changed .ctag{background:#ffedd5;color:#7c2d12;border-color:transparent}
del{color:#b91c1c}ins{color:#15803d;text-decoration:none;font-weight:600}
@media (prefers-color-scheme:dark){del{color:#fca5a5}ins{color:#86efac}}
.nomatch{display:none!important}

footer{max-width:1200px;margin:0 auto;padding:16px 0 28px;border-top:1px solid var(--line);color:var(--muted);font-size:12.5px}
footer p{margin:6px 0}
@media (max-width:520px){.tabs{gap:13px}.tab{font-size:13.5px}.tab .lg{display:none}.tab .sm{display:inline}}
@media print{header.top{position:static}.ctl,.tabs,.toolbars{display:none}.view{display:block!important}}
"""


JS = r"""
(function(){
  'use strict';
  var KEY='hrb-viewer-state-v1';
  function load(){try{return JSON.parse(localStorage.getItem(KEY)||'{}')||{}}catch(e){return {}}}
  function save(s){try{localStorage.setItem(KEY,JSON.stringify(s))}catch(e){}}
  function pad(n){return (n<10?'0':'')+n}
  function isoDate(d){return d.getFullYear()+'-'+pad(d.getMonth()+1)+'-'+pad(d.getDate())}
  function parse(s){if(!s)return null;var m=/^(\d{4})-(\d\d)-(\d\d)T(\d\d):(\d\d)$/.exec(s);return m?new Date(+m[1],m[2]-1,+m[3],+m[4],+m[5]):null}
  var state=load(), now=new Date(), today=isoDate(now);

  // Restore remembered selections; on first visit during the event, jump to today.
  var radios=document.querySelectorAll('input.ctl');
  Array.prototype.forEach.call(radios,function(r){
    if(state[r.name]===r.id) r.checked=true;
    r.addEventListener('change',function(){state[r.name]=r.id;save(state);if(r.name==='view')window.scrollTo(0,0);});
  });
  Array.prototype.forEach.call(document.querySelectorAll('input.ctl[id$="-d-'+today+'"]'),function(r){
    if(!state[r.name]) r.checked=true;
  });
  Array.prototype.forEach.call(document.querySelectorAll('[data-date="'+today+'"]'),function(el){
    if(el.classList.contains('chip')) el.classList.add('today');
  });
  // Close the filter panel after picking a chip.
  Array.prototype.forEach.call(document.querySelectorAll('details.filters'),function(d){
    d.addEventListener('click',function(ev){if(ev.target.closest('label.chip'))setTimeout(function(){d.open=false},120)});
  });
  document.addEventListener('click',function(ev){
    Array.prototype.forEach.call(document.querySelectorAll('details.filters[open]'),function(d){if(!d.contains(ev.target))d.open=false});
  });

  // Freshness: relative time and a warning when the Mac hasn't refreshed the file for a while.
  var checked=document.getElementById('checked');
  function freshness(){
    if(!checked)return;
    var t=new Date(checked.getAttribute('data-iso')), mins=Math.round((Date.now()-t)/60000);
    var rel=mins<1?'just now':mins<60?mins+' min ago':mins<1440?Math.round(mins/60)+' h ago':Math.round(mins/1440)+' d ago';
    checked.textContent=rel;checked.title=t.toLocaleString();
    var stale=document.getElementById('stale');
    if(mins>120){stale.hidden=false;stale.textContent='This copy was last refreshed '+rel+'. The Mac mini may be offline, so tap "Open original" for the newest version.';}
    else stale.hidden=true;
  }

  // Happening now.
  var timed=document.querySelectorAll('[data-s][data-e]'), nowbtn=document.getElementById('nowbtn');
  function tick(){
    var t=new Date(), any=false;
    Array.prototype.forEach.call(timed,function(el){
      var s=parse(el.getAttribute('data-s')), e=parse(el.getAttribute('data-e'));
      var on=!!(s&&e&&s<=t&&t<e);
      el.classList.toggle('now',on);
      if(on&&el.classList.contains('card'))any=true;
      var top=el.querySelector('.c-top');
      if(top){var tag=top.querySelector('.nowtag');if(on&&!tag){tag=document.createElement('span');tag.className='nowtag';tag.textContent='Now';top.appendChild(tag)}else if(!on&&tag)tag.remove()}
    });
    nowbtn.hidden=!any;
    freshness();
  }
  nowbtn.addEventListener('click',function(){
    var el=document.querySelector('.card.now');if(!el)return;
    var view=document.querySelector('input[name=view]');if(view&&!view.checked){view.checked=true;state.view=view.id;save(state)}
    var day=el.closest('.sday'), r=day&&document.getElementById(day.id.replace('-day-','-d-'));
    if(r&&!r.checked){r.checked=true;state[r.name]=r.id;save(state)}
    if(el.offsetParent===null){Array.prototype.forEach.call(document.querySelectorAll('input.ctl[id$="-a-all"],input.ctl[id$="-p-all"]'),function(a){a.checked=true;state[a.name]=a.id});save(state)}
    el.scrollIntoView({behavior:'smooth',block:'center'});
  });

  // Search.
  var bar=document.getElementById('searchbar'), q=document.getElementById('q'), qc=document.getElementById('qcount');
  bar.hidden=false;
  var items=document.querySelectorAll('.card[data-q],.rentry[data-q]'), cells=document.querySelectorAll('.rgrid td[data-q]');
  function search(){
    var terms=q.value.toLowerCase().split(/\s+/).filter(Boolean), n=0;
    function hit(el){var s=el.getAttribute('data-q');return terms.every(function(w){return s.indexOf(w)>=0})}
    Array.prototype.forEach.call(items,function(el){var ok=!terms.length||hit(el);el.classList.toggle('nomatch',!ok);if(ok&&terms.length)n++});
    Array.prototype.forEach.call(cells,function(el){el.classList.toggle('match',!!terms.length&&!!el.getAttribute('data-q')&&hit(el))});
    Array.prototype.forEach.call(document.querySelectorAll('.slot'),function(s){s.classList.toggle('nomatch',!!terms.length&&!s.querySelector('.card:not(.nomatch)'))});
    qc.textContent=terms.length?(n+' match'+(n===1?'':'es')):'';
  }
  q.addEventListener('input',search);

  tick();setInterval(tick,30000);
  document.addEventListener('visibilitychange',function(){if(!document.hidden)tick()});
})();
"""


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(__doc__)
        return 2
    model = load_model(argv[1])
    now = dt.datetime.now().astimezone()
    meta = {"checked_at": now.isoformat(), "changed_at": now.isoformat(), "title": "Schedules + Rehearsals",
            "eyebrow": "HRB 2026 · Producers", "tz": now.strftime("%Z")}
    Path(argv[2]).write_text(render(model, meta), encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
